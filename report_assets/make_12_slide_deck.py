#!/usr/bin/env python3
"""Create a concise 12-slide PS3 presentation from the full report deck.

Build the full deck first with ``python report_assets/make_pptx.py``. This script
keeps the narrative slides below, removes the rest, and renumbers the footer.
It requires ``python-pptx`` (install with ``python -m pip install python-pptx``).
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PS3_Report.pptx"
OUTPUT = ROOT / "PS3_12_Slide_Deck.pptx"

# One-based positions in PS3_Report.pptx. Together these make a short narrative:
# problem and evidence -> approach -> measured results -> safeguards and rollout.
KEEP_SLIDES = (1, 3, 4, 5, 7, 9, 11, 14, 15, 19, 20, 23)
EXPECTED_SOURCE_SLIDES = 23
EXPECTED_OUTPUT_SLIDES = 12


def build() -> Path:
    if not SOURCE.is_file():
        raise FileNotFoundError(
            f"Missing source deck: {SOURCE}. Build it first with report_assets/make_pptx.py."
        )

    presentation = Presentation(SOURCE)
    if len(presentation.slides) != EXPECTED_SOURCE_SLIDES:
        raise ValueError(
            f"Expected {EXPECTED_SOURCE_SLIDES} source slides, "
            f"found {len(presentation.slides)}. Review KEEP_SLIDES before rebuilding."
        )

    slide_id_list = presentation.slides._sldIdLst
    original_ids = list(slide_id_list)
    keep = set(KEEP_SLIDES)
    for position, slide_id in reversed(list(enumerate(original_ids, start=1))):
        if position not in keep:
            slide_id_list.remove(slide_id)
            presentation.part.drop_rel(slide_id.rId)

    if len(presentation.slides) != EXPECTED_OUTPUT_SLIDES:
        raise AssertionError(f"Expected 12 slides after selection, got {len(presentation.slides)}")

    # The source deck has no number on its title slide. Renumber remaining page
    # markers sequentially without replacing their styled text runs.
    for page_number, slide in enumerate(presentation.slides, start=1):
        if page_number == 1:
            continue
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.left < Inches(11.5) or shape.top < Inches(7.0):
                continue
            paragraphs = shape.text_frame.paragraphs
            if paragraphs and paragraphs[0].runs:
                paragraphs[0].runs[0].text = str(page_number)
                break

    presentation.save(OUTPUT)

    # Re-open the generated package to catch relationship/serialization issues.
    check = Presentation(OUTPUT)
    if len(check.slides) != EXPECTED_OUTPUT_SLIDES:
        raise AssertionError(f"Saved deck has {len(check.slides)} slides, not 12")
    print(f"Wrote {OUTPUT.relative_to(ROOT)} ({len(check.slides)} slides)")
    return OUTPUT


if __name__ == "__main__":
    build()
